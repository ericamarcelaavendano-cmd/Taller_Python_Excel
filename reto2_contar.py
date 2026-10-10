# ==============================================================
#  RETO 2: Mostrar cuántos registros tiene el archivo
#  Estudiante: Erica Avendaño
# ==============================================================
import openpyxl

libro = openpyxl.load_workbook("personas.xlsx")
hoja = libro.active

# Guardamos todas las personas en una lista (sin el encabezado)
personas = list(hoja.iter_rows(min_row=2, values_only=True))

# len() nos dice cuántos elementos tiene la lista
total = len(personas)

print("=" * 45)
print(f"  El archivo personas.xlsx tiene {total} registros")
print("=" * 45)

# Dato curioso: max_row cuenta también la fila del encabezado
print(f"(Filas totales en Excel: {hoja.max_row}  ->  1 encabezado + {total} personas)")
