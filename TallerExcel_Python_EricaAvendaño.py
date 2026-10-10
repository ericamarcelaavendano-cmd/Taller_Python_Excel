
#  TALLER PYTHON – Lectura de archivos Excel

# PASO 1. Importar librerías
import openpyxl                      # Para leer archivos de Excel (.xlsx)
from tabulate import tabulate        # Para mostrar los datos como tabla en consola
from collections import Counter      # (Parte 2) Para contar cuántas veces se repite algo
import matplotlib.pyplot as plt      # (Parte 2) Para dibujar las gráficas


# PASO 2. Abrir el archivo Excel
excel_dataframe = openpyxl.load_workbook("personas.xlsx")

# PASO 3. Seleccionar la hoja excel activa 
dataframe = excel_dataframe.active


# PASO 4. Crear una lista vacía donde guardaremos toda la información
data = []

# PASO 5. Recorrer las filas
# range(1, max_row) genera 1, 2, ..., 35 una vuelta por cada persona.
for row in range(1, dataframe.max_row):

# PASO 6. Crear la fila temporal
    _row = [row]


    # PASO 7. Recorrer las columnas
    # iter_cols() entrega cada columna completa (desde la fila 1 del Excel).
    # col[0] es el encabezado, así que col[row] con row >= 1 salta el
    # encabezado y toma el dato de la persona actual.
    # ----------------------------------------------------------
    for col in dataframe.iter_cols(1, dataframe.max_column):
        _row.append(col[row].value)

    # PASO 8. Guardar la fila completa en la lista principal
    data.append(_row)

# PASO 9. Crear encabezados y su alineación
headers = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
headers_align = ("center",) * 6

# PASO 10. Mostrar el resultado en una tabla organizada 
print(tabulate(data, headers=headers, tablefmt="fancy_grid", colalign=headers_align))


 # Como soy un poco curiosa decidí investigar y crear una grafica como informe visual 
 #Donde se puede visualizar el informe detalladamente. 
 
def tipo_empresa(nombre):
    """Devuelve el tipo de empresa según cómo está escrito su nombre."""
    if nombre.endswith("and Sons"):
        return "Familiar (Evans and Sons)"
    if " and " in nombre:
        return "Tres socios (Adkins, Stewart and Tran)"
    for sufijo in ["Group", "Inc", "PLC", "LLC", "Ltd"]:
        if nombre.endswith(sufijo):
            return sufijo
    return "Dos apellidos (Smith-Holmes)"   # Las demás tienen guion


# 1. Contar cuántas hay de cada categoría
total = len(data)
empresas = Counter(tipo_empresa(fila[3]) for fila in data)
dominios = Counter(fila[4].split("@")[1] for fila in data)

# 2. Mostrar el resumen en consola
print(tabulate(empresas.most_common(), headers=["Tipo de empresa", "Cantidad"], tablefmt="fancy_grid"))
print(tabulate(dominios.most_common(), headers=["Dominio", "Cantidad"], tablefmt="fancy_grid"))

# 3. Crear la figura con 2 gráficas lado a lado
fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(13, 5))
fig.suptitle(f"Análisis de personas·  {total} registros", fontsize=15, fontweight="bold")

# 4. Dibujar las dos gráficas con el mismo ciclo
for ax, conteo, titulo in [(ax1, empresas, "¿Qué tipo de empresa es?"),
                           (ax2, dominios, "Dominio del correo electrónico")]:
    etiquetas, valores = zip(*reversed(conteo.most_common()))   # la mayor queda arriba
    barras = ax.barh(etiquetas, valores, color="#9D3BCB", height=0.6)
    ax.bar_label(barras, labels=[f"{v} ({v / total:.0%})" for v in valores], padding=4)
    ax.set_title(titulo, loc="left", fontweight="bold")
    ax.set_xlim(0, max(valores) * 1.3)                       # espacio para las etiquetas
    ax.set_ylim(len(valores) - 8.6, len(valores) - 0.4)      # barras del mismo grosor
    ax.set_xticks([])                                        # quitamos los números de abajo
    ax.spines[["top", "right", "bottom", "left"]].set_visible(False)   # quitamos los bordes

# 5. Guardar la imagen y mostrar la ventana
plt.tight_layout()
plt.savefig("graficas_personas.png", dpi=150)
print("Gráfica guardada como graficas_personas.png")
plt.show()
